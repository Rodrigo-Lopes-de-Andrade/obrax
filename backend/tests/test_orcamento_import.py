"""B5.5 — importação XLSX da EAP (US-13): parsing, validação linha-a-linha,
criação de hierarquia, template de download e round-trip.
"""

from __future__ import annotations

from collections.abc import Callable
from io import BytesIO

from openpyxl import Workbook, load_workbook
from starlette.testclient import TestClient

from app.modules.orcamento.importers import xlsx_importer
from tests.conftest import SeededUser

SeedTenant = Callable[..., SeededUser]

CNPJ_VALIDO = "11222333000181"
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _seed_versao(client: TestClient, owner: SeededUser) -> str:
    cli = client.post(
        "/api/v1/clientes",
        headers=owner.auth_header,
        json={"tipo_pessoa": "PJ", "documento": CNPJ_VALIDO, "nome": "Cliente"},
    ).json()
    obra = client.post(
        "/api/v1/obras",
        headers=owner.auth_header,
        json={"cliente_id": cli["id"], "codigo": "OBR-1", "nome": "Obra"},
    ).json()
    orc = client.post(
        f"/api/v1/obras/{obra['id']}/orcamentos",
        headers=owner.auth_header,
        json={"titulo": "Orç"},
    ).json()
    return str(orc["versao"]["id"])


def _xlsx(rows: list[list[object]], *, header: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header or xlsx_importer.HEADERS)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client: TestClient, owner: SeededUser, vid: str, content: bytes) -> object:
    return client.post(
        f"/api/v1/orcamentos/versoes/{vid}/importar-planilha",
        headers=owner.auth_header,
        files={"arquivo": ("orc.xlsx", content, _XLSX_MEDIA)},
    )


def _por_tipo(tree: dict, tipo: str) -> list[dict]:
    return [i for i in tree["itens"] if i["tipo"] == tipo]


def test_import_valido_cria_hierarquia(client: TestClient, seed_tenant: SeedTenant) -> None:
    owner = seed_tenant(role="owner")
    vid = _seed_versao(client, owner)
    content = _xlsx(
        [
            ["F", "", "Fundações", "", "", ""],
            ["S", "", "Sapatas", "", "", ""],
            ["I", "SINAPI-1", "Concreto usinado", "m³", 10, 100.00],
            ["I", "SINAPI-2", "Armadura", "kg", 50, 12.00],
            ["S", "", "Vigas", "", "", ""],
            ["I", "SINAPI-3", "Fôrma", "m²", 20, 30.00],
        ]
    )
    resp = _upload(client, owner, vid, content)
    assert resp.status_code == 201, resp.text
    tree = resp.json()
    assert len(_por_tipo(tree, "F")) == 1
    assert len(_por_tipo(tree, "S")) == 2
    assert len(_por_tipo(tree, "I")) == 3
    # 10×100 + 50×12 + 20×30 = 1000 + 600 + 600 = 2200
    assert tree["versao"]["custo_direto"] == "2200.00"
    assert tree["versao"]["preco_venda"] == "2200.00"


def test_import_erros_linha_a_linha_422(client: TestClient, seed_tenant: SeedTenant) -> None:
    owner = seed_tenant(role="owner")
    vid = _seed_versao(client, owner)
    content = _xlsx(
        [
            ["S", "", "Serviço órfão", "", "", ""],  # linha 2: serviço sem fase
            ["F", "", "Fase OK", "", "", ""],
            ["I", "", "", "", "abc", ""],  # linha 4: insumo sem serviço + descrição + qtd inválida
            ["Z", "", "Tipo inválido", "", "", ""],  # linha 5: tipo inválido
        ]
    )
    resp = _upload(client, owner, vid, content)
    assert resp.status_code == 422
    erros = resp.json()["detail"]["erros"]
    linhas_campos = {(e["linha"], e["campo"]) for e in erros}
    assert (2, "tipo") in linhas_campos  # serviço sem fase
    assert (4, "descricao") in linhas_campos
    assert (4, "quantidade") in linhas_campos
    assert (5, "tipo") in linhas_campos
    # Nada foi criado (import é tudo-ou-nada).
    tree = client.get(f"/api/v1/orcamentos/versoes/{vid}/itens", headers=owner.auth_header).json()
    assert tree["itens"] == []


def test_import_cabecalho_incompleto_422(client: TestClient, seed_tenant: SeedTenant) -> None:
    owner = seed_tenant(role="owner")
    vid = _seed_versao(client, owner)
    # falta a coluna 'descricao'
    content = _xlsx([["F", "x"]], header=["tipo", "codigo"])
    resp = _upload(client, owner, vid, content)
    assert resp.status_code == 422
    assert any(e["campo"] == "descricao" for e in resp.json()["detail"]["erros"])


def test_import_arquivo_invalido_422(client: TestClient, seed_tenant: SeedTenant) -> None:
    owner = seed_tenant(role="owner")
    vid = _seed_versao(client, owner)
    resp = _upload(client, owner, vid, b"isto nao e um xlsx")
    assert resp.status_code == 422


def test_import_versao_travada_409(client: TestClient, seed_tenant: SeedTenant) -> None:
    owner = seed_tenant(role="owner")
    vid = _seed_versao(client, owner)
    orc_id = client.get("/api/v1/orcamentos", headers=owner.auth_header).json()["items"][0]["id"]
    client.patch(
        f"/api/v1/orcamentos/{orc_id}/status",
        headers=owner.auth_header,
        json={"status": "aguardando_aprovacao"},
    )
    client.patch(
        f"/api/v1/orcamentos/{orc_id}/status",
        headers=owner.auth_header,
        json={"status": "aprovado"},
    )
    content = _xlsx([["F", "", "Fase", "", "", ""]])
    assert _upload(client, owner, vid, content).status_code == 409


def test_template_download_e_round_trip(client: TestClient, seed_tenant: SeedTenant) -> None:
    owner = seed_tenant(role="owner")
    vid = _seed_versao(client, owner)

    resp = client.get("/api/v1/orcamentos/importar/modelo", headers=owner.auth_header)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == _XLSX_MEDIA
    # É um XLSX legível com o cabeçalho esperado.
    wb = load_workbook(BytesIO(resp.content))
    assert [c.value for c in wb.active[1]] == xlsx_importer.HEADERS

    # O próprio modelo (com linhas de exemplo) importa sem erros.
    up = _upload(client, owner, vid, resp.content)
    assert up.status_code == 201, up.text
    assert len(up.json()["itens"]) == 6  # 1 fase, 2 serviços, 3 insumos


def test_parse_excede_5000_linhas() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(xlsx_importer.HEADERS)
    for _ in range(5001):
        ws.append(["F", "", "Fase", "", "", ""])
    buf = BytesIO()
    wb.save(buf)
    parsed = xlsx_importer.parse_xlsx(buf.getvalue())
    assert parsed.erros
    assert "excede" in parsed.erros[0]["mensagem"]
